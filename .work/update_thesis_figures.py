from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


SOURCE = Path(r"C:\xampp\htdocs\qr-code-attendance-system-main\.work\Thesis-source.docx")
IPO_IMAGE = Path(r"C:\Users\ADMIN\AppData\Local\Temp\codex-clipboard-2ae641da-4f56-41a2-be4a-e63cece9a0e0.png")
AGILE_IMAGE = Path(r"C:\Users\ADMIN\AppData\Local\Temp\codex-clipboard-94576b27-5b3f-4a4e-a948-3f6f92e08354.png")
OUTPUT = Path(r"C:\xampp\htdocs\qr-code-attendance-system-main\deliverables\Chapters 1 and 2 - Revised with IPO and Agile.docx")


def clear_paragraph_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_alt_text(paragraph, title, description):
    doc_prs = paragraph._p.xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[0].set("title", title)
        doc_prs[0].set("descr", description)


def insert_paragraph_after(paragraph, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    return new_paragraph


def set_body_paragraph(paragraph, lead, body):
    clear_paragraph_content(paragraph)
    lead_run = paragraph.add_run(lead)
    lead_run.bold = True
    paragraph.add_run(body)


doc = Document(SOURCE)

# Keep the manually typed table of contents synchronized with the revised pagination.
doc.paragraphs[61].runs[0].text = "Locale of the Study\t32"
doc.paragraphs[62].runs[0].text = "Respondents\t33"

# Replace the existing conceptual framework visual with the user-supplied IPO figure.
ipo_paragraph = doc.paragraphs[133]
clear_paragraph_content(ipo_paragraph)
ipo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
ipo_paragraph.paragraph_format.keep_with_next = True
ipo_paragraph.add_run().add_picture(str(IPO_IMAGE), width=Inches(6.05))
set_alt_text(
    ipo_paragraph,
    "Input Process Output conceptual framework",
    "Input, process, and output framework for the Smart NSTP Management and AI-Integrated Learning Platform.",
)

# Align the methodology text with the six Agile phases shown in the supplied diagram.
doc.paragraphs[210].text = (
    "The study will apply the Agile Software Development Methodology through six iterative phases: "
    "Analysis, Design, Develop, Testing, Implement, and Review and Refinement. Each phase will produce "
    "reviewable outputs and may return to an earlier phase when requirements, defects, or risks require "
    "correction. A version-controlled repository and change log will preserve traceability between "
    "stakeholder needs, implemented features, test cases, defects, and revisions. Figure 2 presents the "
    "development cycle used in the study."
)
doc.paragraphs[210].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

set_body_paragraph(
    doc.paragraphs[211],
    "Analysis. ",
    "The researcher will document current NSTP workflows, records, reports, learning activities, pain "
    "points, and control requirements through interviews, observation, and review of authorized forms or "
    "reports. Stakeholder needs will be translated into functional and non-functional requirements, user "
    "roles, permissions, business rules, report requirements, acceptance criteria, and an initial product "
    "backlog. Privacy, security, accessibility, data quality, backup, audit logging, and human-review "
    "requirements will be defined before implementation.",
)
set_body_paragraph(
    doc.paragraphs[212],
    "Design. ",
    "The researcher will prepare the database schema, process models, role-permission matrix, system "
    "architecture, navigation structure, wireframes, report layouts, and testable interface designs. The "
    "design will separate presentation, application, data, and AI-service concerns where practical to "
    "improve maintainability and control access. Representative users will review prototypes for clarity, "
    "task fit, terminology, and missing functions before full development.",
)
set_body_paragraph(
    doc.paragraphs[213],
    "Develop. ",
    "Core administrative features will be implemented first, followed by learning, assessment, reporting, "
    "communication, QR, and AI-assisted functions. Each working increment will undergo developer checks "
    "before integration. AI functions will connect only to approved information sources and defined tasks. "
    "Prompts, outputs, fallback messages, and escalation paths will keep suggestions reviewable and prevent "
    "them from becoming autonomous official decisions.",
)
set_body_paragraph(
    doc.paragraphs[214],
    "Testing. ",
    "Unit, integration, system, permission, data-validation, report, backup-recovery, and acceptance tests "
    "will be executed against documented requirements. Critical workflows will include registration, "
    "section assignment, attendance capture and correction, score and grade handling, content access, "
    "report generation, and account authorization. Defects will be classified by severity, corrected, and "
    "retested before implementation proceeds.",
)
set_body_paragraph(
    doc.paragraphs[215],
    "Implement. ",
    "A stable build will be deployed in a controlled TAU-NSTP setting using representative accounts and "
    "approved data. Eligible participants will complete role-appropriate tasks while the researcher records "
    "task outcomes, technical observations, and operational issues. No controlled pilot will proceed while "
    "a known critical defect could expose records, corrupt official data, bypass authorization, or produce "
    "an irreversible academic decision.",
)
set_body_paragraph(
    doc.paragraphs[216],
    "Review and Refinement. ",
    "Technical findings, user ratings, qualitative feedback, and operational logs will be reviewed together. "
    "Agreed corrections will be added to the backlog, implemented in the next iteration, and regression-tested. "
    "The cycle will continue until the evaluated release satisfies the approved requirements and the researcher "
    "can prepare final implementation recommendations.",
)

for paragraph in doc.paragraphs[211:217]:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Add the Agile figure and captions immediately after the six-phase explanation.
agile_paragraph = insert_paragraph_after(doc.paragraphs[216], style=doc.paragraphs[133].style)
agile_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
agile_paragraph.paragraph_format.keep_with_next = True
agile_paragraph.add_run().add_picture(str(AGILE_IMAGE), width=Inches(5.35))
set_alt_text(
    agile_paragraph,
    "Agile development cycle",
    "Six-phase Agile cycle consisting of Analysis, Design, Develop, Testing, Implement, and Review and Refinement.",
)

figure_number = insert_paragraph_after(agile_paragraph, style=doc.paragraphs[134].style)
figure_number.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure_number.paragraph_format.space_before = doc.paragraphs[134].paragraph_format.space_before
figure_number.paragraph_format.keep_with_next = True
number_run = figure_number.add_run("Figure 2")
number_run.bold = True
number_run.font.name = doc.paragraphs[134].runs[0].font.name
number_run.font.size = doc.paragraphs[134].runs[0].font.size

figure_title = insert_paragraph_after(figure_number, style=doc.paragraphs[135].style)
figure_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure_title.paragraph_format.space_after = doc.paragraphs[135].paragraph_format.space_after
title_run = figure_title.add_run("Agile Software Development Cycle of the Study")
title_run.italic = True
title_run.font.name = doc.paragraphs[135].runs[0].font.name
title_run.font.size = doc.paragraphs[135].runs[0].font.size

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
